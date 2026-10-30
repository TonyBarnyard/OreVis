"""Office-document outputs: live-formula Excel workbooks, Word credit memos, charts.

Workbooks are written with formulas (not pasted values) so the broker can
change an assumption and watch every ratio update -- the same expectation a
credit committee has of an analyst's spreadsheet.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from docx import Document
from docx.shared import Inches

from . import underwriting as uw

FONT = "Arial"
INPUT = Font(name=FONT, color="0000FF")
BOLD = Font(name=FONT, bold=True)
BODY = Font(name=FONT)
KEY = PatternFill("solid", start_color="FFFF00")

DISCLAIMER = ("Prepared by an AI analysis assistant for review by a qualified lending "
              "professional. This is not a credit decision, a commitment to lend or arrange "
              "financing, or legal advice. All figures must be verified against source "
              "documents.")


def build_loan_sizing_workbook(path: str | Path, deal_name: str, noi: float,
                               collateral_value: float, loan_amount: float,
                               annual_rate: float, amort_years: float, term_years: float,
                               min_dscr: float = 1.25, max_ltv: float = 0.75,
                               min_debt_yield: float = 0.09) -> Path:
    path = Path(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sizing"
    ws["A1"] = f"Loan Sizing - {deal_name}"
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws["A2"] = "Blue = input. Change any input; all metrics recalculate."
    ws["A2"].font = Font(name=FONT, italic=True)

    inputs = [
        ("Net operating income ($)", noi, "$#,##0"),
        ("Collateral value ($)", collateral_value, "$#,##0"),
        ("Requested loan ($)", loan_amount, "$#,##0"),
        ("Interest rate", annual_rate, "0.000%"),
        ("Amortization (years)", amort_years, "0"),
        ("Term (years)", term_years, "0"),
        ("Policy min DSCR", min_dscr, "0.00x"),
        ("Policy max LTV", max_ltv, "0.0%"),
        ("Policy min debt yield", min_debt_yield, "0.0%"),
    ]
    for i, (label, value, fmt) in enumerate(inputs, start=4):
        ws.cell(i, 1, label).font = BODY
        c = ws.cell(i, 2, value)
        c.font, c.number_format = INPUT, fmt
    # B4 NOI, B5 value, B6 loan, B7 rate, B8 amort, B9 term, B10 dscr, B11 ltv, B12 dy

    outputs = [
        ("Annual debt service ($)", "=-PMT(B7/12,B8*12,B6)*12", "$#,##0"),
        ("DSCR", "=IFERROR(B4/B15,0)", "0.00x"),
        ("LTV", "=IFERROR(B6/B5,0)", "0.0%"),
        ("Debt yield", "=IFERROR(B4/B6,0)", "0.0%"),
        ("Max loan @ min DSCR ($)", "=IFERROR(-PV(B7/12,B8*12,B4/B10/12),0)", "$#,##0"),
        ("Max loan @ max LTV ($)", "=B5*B11", "$#,##0"),
        ("Max loan @ min debt yield ($)", "=IFERROR(B4/B12,0)", "$#,##0"),
        ("Max supportable loan ($)", "=MIN(B19:B21)", "$#,##0"),
        ("Balloon balance at maturity ($)", "=-FV(B7/12,B9*12,PMT(B7/12,B8*12,B6),B6)", "$#,##0"),
        ("Meets policy?", '=IF(AND(B16>=B10,B17<=B11,B18>=B12),"Yes","Exception - review")', "@"),
    ]
    ws["A14"], ws["A14"].font = "Results", BOLD
    for i, (label, formula, fmt) in enumerate(outputs, start=15):
        ws.cell(i, 1, label).font = BODY
        c = ws.cell(i, 2, formula)
        c.font, c.number_format = BODY, fmt
    ws["B22"].fill = KEY

    ws["A27"] = DISCLAIMER
    ws["A27"].font = Font(name=FONT, italic=True, size=8)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18

    # Rate-shock sheet with formulas tied to the Sizing inputs, plus a chart.
    st = wb.create_sheet("Rate Shock")
    for col, h in enumerate(["Shock (bps)", "Rate", "Debt service ($)", "DSCR"], start=1):
        st.cell(1, col, h).font = BOLD
    for r, bps in enumerate((0, 50, 100, 150, 200, 300), start=2):
        st.cell(r, 1, bps).font = INPUT
        st.cell(r, 2, f"=Sizing!$B$7+A{r}/10000").number_format = "0.000%"
        st.cell(r, 3, f"=-PMT(B{r}/12,Sizing!$B$8*12,Sizing!$B$6)*12").number_format = "$#,##0"
        st.cell(r, 4, f"=IFERROR(Sizing!$B$4/C{r},0)").number_format = "0.00x"
    chart = LineChart()
    chart.title, chart.y_axis.title, chart.x_axis.title = "DSCR under rate shock", "DSCR", "bps"
    chart.add_data(Reference(st, min_col=4, min_row=1, max_row=7), titles_from_data=True)
    chart.set_categories(Reference(st, min_col=1, min_row=2, max_row=7))
    st.add_chart(chart, "F2")

    # Amortization schedule through term.
    sa = wb.create_sheet("Amortization")
    for col, h in enumerate(["Month", "Payment", "Interest", "Principal", "Balance"], start=1):
        sa.cell(1, col, h).font = BOLD
    for row in uw.amortization_schedule(loan_amount, annual_rate, amort_years, term_years):
        sa.append([row["month"], row["payment"], row["interest"], row["principal"], row["balance"]])
    bar = BarChart()
    bar.title = "Loan balance"
    bar.add_data(Reference(sa, min_col=5, min_row=1, max_row=sa.max_row), titles_from_data=True)
    sa.add_chart(bar, "G2")

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def build_spreadsheet(path: str | Path, sheets: dict[str, list[list]]) -> Path:
    """General workbook: {sheet name: rows}. Strings starting with "=" are formulas."""
    path = Path(path)
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(str(name)[:31])
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(r, c, value)
                cell.font = BOLD if r == 1 else BODY
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = "#,##0.00" if abs(value) < 10 else "#,##0"
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def build_credit_memo(path: str | Path, deal_name: str, sections: dict[str, str],
                      metrics: dict[str, str] | None = None,
                      chart_png: str | Path | None = None) -> Path:
    """Word credit memo. `sections` maps heading -> body text, in order."""
    path = Path(path)
    doc = Document()
    doc.add_heading(f"Credit Memorandum - {deal_name}", 0)
    p = doc.add_paragraph(DISCLAIMER)
    p.runs[0].italic = True
    if metrics:
        doc.add_heading("Key metrics", 1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Metric", "Value"
        for k, v in metrics.items():
            cells = table.add_row().cells
            cells[0].text, cells[1].text = k, v
    for heading, body in sections.items():
        doc.add_heading(heading, 1)
        for para in body.split("\n\n"):
            doc.add_paragraph(para.strip())
    if chart_png:
        doc.add_picture(str(chart_png), width=Inches(6))
    doc.add_heading("Broker review", 1)
    doc.add_paragraph("Reviewed and approved for submission by: ____________   "
                      "Date: ________")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path


def build_chart(path: str | Path, title: str, labels: list[str], values: list[float],
                kind: str = "bar", y_label: str = "") -> Path:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    path = Path(path)
    fig, ax = plt.subplots(figsize=(8, 4.5))
    if kind == "line":
        ax.plot(labels, values, marker="o", color="#1f4e79")
    else:
        ax.bar(labels, values, color="#1f4e79")
    ax.set_title(title)
    ax.set_ylabel(y_label)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=150)
    plt.close(fig)
    return path
