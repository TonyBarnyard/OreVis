from openpyxl import load_workbook
from docx import Document

from lending_agent import documents
from lending_agent.agent import build_tools


def test_workbook_has_live_formulas(tmp_path):
    p = documents.build_loan_sizing_workbook(tmp_path / "deal.xlsx", "Test Plaza", 200_000,
                                             3_000_000, 2_000_000, 0.07, 25, 10)
    ws = load_workbook(p)["Sizing"]
    assert ws["B16"].value == "=IFERROR(B4/B15,0)"
    assert ws["B22"].value == "=MIN(B19:B21)"
    assert {"Sizing", "Rate Shock", "Amortization"} <= set(load_workbook(p).sheetnames)


def test_credit_memo_contains_sections_and_disclaimer(tmp_path):
    p = documents.build_credit_memo(tmp_path / "memo.docx", "Test Plaza",
                                    {"Executive Summary": "Stabilized retail."},
                                    {"DSCR": "1.32x"})
    text = "\n".join(par.text for par in Document(p).paragraphs)
    assert "Executive Summary" in text and "not a credit decision" in text


def test_tools_cannot_escape_workspace(tmp_path):
    tools = {t.name: t for t in build_tools(tmp_path / "tenant-a")}
    try:
        tools["create_chart"].call({"filename": "../tenant-b/x.png", "title": "t",
                                    "labels": ["a"], "values": [1]})
    except ValueError:
        return
    raise AssertionError("path traversal was not blocked")
